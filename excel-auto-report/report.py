"""把多間門市格式不一的銷售 Excel，一鍵合併、清理並產出月報。

用法：
    python report.py data/*.xlsx                 # 自動選最新月份
    python report.py data/*.xlsx --month 2026-08 -o output/八月月報.xlsx
"""
from __future__ import annotations

import argparse
import glob
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 各門市習慣的欄位名稱 → 標準欄位
COLUMN_ALIASES = {
    "單號": ["單號", "訂單編號", "交易序號", "order_id"],
    "日期": ["日期", "交易日期", "銷售日期", "date"],
    "門市": ["門市", "店名", "分店", "store"],
    "商品": ["商品", "商品名稱", "品名", "品項", "product"],
    "類別": ["類別", "分類", "category"],
    "數量": ["數量", "銷售數量", "qty", "quantity"],
    "單價": ["單價", "售價", "價格", "price"],
}
REQUIRED = ["日期", "商品", "數量", "單價"]


@dataclass
class CleanLog:
    """記錄清理過程，最後輸出成「清理紀錄」工作表，讓客戶知道資料被怎麼處理。"""

    entries: list[dict] = field(default_factory=list)

    def add(self, store: str, problem: str, count: int, action: str):
        if count:
            self.entries.append({"門市": store, "問題": problem, "筆數": count, "處理方式": action})

    def to_frame(self) -> pd.DataFrame:
        return pd.DataFrame(self.entries, columns=["門市", "問題", "筆數", "處理方式"])


def normalize_text(value) -> str:
    """全形轉半形、去掉頭尾與重複空白：「 咖啡豆（半磅） 」→「咖啡豆(半磅)」。"""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", " ", text).strip()


def standardize_columns(df: pd.DataFrame, source: str) -> pd.DataFrame:
    lookup = {alias.lower(): std for std, aliases in COLUMN_ALIASES.items() for alias in aliases}
    df = df.rename(columns={c: lookup.get(normalize_text(c).lower(), c) for c in df.columns})
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        raise ValueError(f"{source} 缺少必要欄位：{', '.join(missing)}（現有欄位：{list(df.columns)}）")
    return df


def load_file(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path) if path.suffix.lower() in {".xlsx", ".xls"} else pd.read_csv(path)
    df = standardize_columns(df, path.name)
    if "門市" not in df.columns:
        # 檔名「台北店_2026Q3.xlsx」→ 門市「台北店」
        df["門市"] = path.stem.split("_")[0]
    if "類別" not in df.columns:
        df["類別"] = "未分類"
    if "單號" not in df.columns:
        df["單號"] = ""  # 沒有單號時，只能用「所有欄位都相同」判斷重複
    return df[["單號", "日期", "門市", "商品", "類別", "數量", "單價"]]


def clean(raw: pd.DataFrame) -> tuple[pd.DataFrame, CleanLog]:
    log = CleanLog()
    parts = []
    for store, df in raw.groupby("門市", sort=False):
        df = df.copy()

        df["商品"] = df["商品"].map(normalize_text)
        df["類別"] = df["類別"].fillna("未分類").map(normalize_text)

        df["日期"] = pd.to_datetime(df["日期"].astype(str).str.replace("/", "-"), errors="coerce", format="mixed")
        bad_date = df["日期"].isna()
        log.add(store, "日期無法辨識", int(bad_date.sum()), "刪除該列")
        df = df[~bad_date]

        qty_text = df["數量"].astype(str).str.extract(r"(-?\d+)")[0]
        df["數量"] = pd.to_numeric(qty_text, errors="coerce")
        bad_qty = df["數量"].isna() | (df["數量"] <= 0)
        log.add(store, "數量無法辨識或為 0", int(bad_qty.sum()), "刪除該列")
        df = df[~bad_qty]

        df["單價"] = pd.to_numeric(df["單價"], errors="coerce")
        no_price = df["單價"].isna()
        df["單價"] = df["單價"].fillna(df.groupby("商品")["單價"].transform("median"))
        log.add(store, "單價空白", int(no_price.sum()), "以同商品的中位數單價補上")
        still_missing = df["單價"].isna()
        log.add(store, "單價空白且無同商品可參考", int(still_missing.sum()), "刪除該列")
        df = df[~still_missing]

        dup = df.duplicated()
        log.add(store, "完全重複的列", int(dup.sum()), "只保留一筆")
        df = df[~dup]

        parts.append(df)

    data = pd.concat(parts, ignore_index=True)
    data["數量"] = data["數量"].astype(int)
    data["金額"] = data["數量"] * data["單價"]
    return data.sort_values(["日期", "門市"]).reset_index(drop=True), log


# ---------- 報表輸出 ----------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F4E78")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0'
PCT = '+0.0%;-0.0%;0.0%'


def write_table(ws, df: pd.DataFrame, start_row: int = 1, formats: dict[str, str] | None = None):
    formats = formats or {}
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=col)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=1):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime().date()
            cell = ws.cell(row=i, column=j, value=value)
            cell.border = BORDER
            fmt = formats.get(df.columns[j - 1])
            if fmt:
                cell.number_format = fmt
    for j, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df.iloc[:200, j - 1]]) * 1.8 + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 10), 40)
    return start_row + len(df)


def growth(current: float, previous: float) -> float | None:
    return None if not previous else current / previous - 1


def build_report(data: pd.DataFrame, log: CleanLog, month: str, out: Path) -> Path:
    period = pd.Period(month, freq="M")
    data = data.assign(月份=data["日期"].dt.to_period("M"))
    cur, prev = data[data["月份"] == period], data[data["月份"] == period - 1]
    if cur.empty:
        raise ValueError(f"資料中沒有 {month} 的銷售紀錄，可用月份：{sorted(map(str, data['月份'].unique()))}")

    wb = Workbook()

    # 摘要
    ws = wb.active
    ws.title = "摘要"
    ws["A1"] = f"{period.year} 年 {period.month} 月 銷售月報"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"資料來源：{', '.join(data['門市'].unique())}　|　由 report.py 自動產生"
    ws["A2"].font = Font(color="808080")
    kpis = pd.DataFrame(
        [
            ("總營收", cur["金額"].sum(), prev["金額"].sum()),
            ("交易筆數", len(cur), len(prev)),
            ("銷售件數", cur["數量"].sum(), prev["數量"].sum()),
            ("平均每筆金額", cur["金額"].mean(), prev["金額"].mean() if len(prev) else 0),
        ],
        columns=["指標", "本月", "上月"],
    )
    kpis["月增率"] = [growth(c, p) for c, p in zip(kpis["本月"], kpis["上月"])]
    write_table(ws, kpis, start_row=4, formats={"本月": MONEY, "上月": MONEY, "月增率": PCT})
    best_store = cur.groupby("門市")["金額"].sum().idxmax()
    best_item = cur.groupby("商品")["金額"].sum().idxmax()
    ws["A10"] = f"本月業績最高門市：{best_store}；營收最高商品：{best_item}"
    ws["A10"].font = Font(bold=True)

    # 門市比較
    ws = wb.create_sheet("門市比較")
    stores = cur.groupby("門市").agg(營收=("金額", "sum"), 交易筆數=("金額", "size"), 銷售件數=("數量", "sum"))
    stores["上月營收"] = prev.groupby("門市")["金額"].sum().reindex(stores.index).fillna(0)
    stores["月增率"] = [growth(c, p) for c, p in zip(stores["營收"], stores["上月營收"])]
    stores["營收占比"] = stores["營收"] / stores["營收"].sum()
    stores = stores.sort_values("營收", ascending=False).reset_index()
    last = write_table(ws, stores, formats={"營收": MONEY, "上月營收": MONEY, "月增率": PCT, "營收占比": "0.0%"})
    chart = BarChart()
    chart.title, chart.y_axis.title = "各門市營收", "營收 (元)"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "I2")

    # 商品排行
    ws = wb.create_sheet("商品排行")
    items = (
        cur.groupby(["商品", "類別"])
        .agg(銷售件數=("數量", "sum"), 營收=("金額", "sum"))
        .sort_values("營收", ascending=False)
        .reset_index()
    )
    items.insert(0, "排名", range(1, len(items) + 1))
    last = write_table(ws, items, formats={"營收": MONEY})
    chart = BarChart()
    chart.type, chart.title = "bar", "商品營收排行"
    chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=min(last, 11)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=2, max_row=min(last, 11)))
    chart.y_axis.scaling.orientation = "minMax"
    chart.x_axis.scaling.orientation = "maxMin"
    chart.height, chart.width = 9, 16
    ws.add_chart(chart, "H2")

    # 每日趨勢（各門市一條線）
    ws = wb.create_sheet("每日趨勢")
    daily = cur.pivot_table(index=cur["日期"].dt.date, columns="門市", values="金額", aggfunc="sum", fill_value=0)
    daily["合計"] = daily.sum(axis=1)
    daily = daily.reset_index()
    daily.columns.name = None
    last = write_table(ws, daily, formats={c: MONEY for c in daily.columns[1:]})
    for row in ws.iter_rows(min_row=2, max_row=last, max_col=1):
        row[0].number_format = "mm/dd"
    chart = LineChart()
    chart.title, chart.y_axis.title = "每日營收", "營收 (元)"
    chart.add_data(Reference(ws, min_col=2, max_col=len(daily.columns) - 1, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    chart.height, chart.width = 9, 22
    ws.add_chart(chart, f"{get_column_letter(len(daily.columns) + 2)}2")

    # 清理紀錄
    ws = wb.create_sheet("清理紀錄")
    log_df = log.to_frame()
    if log_df.empty:
        ws["A1"] = "資料沒有發現需要清理的問題。"
    else:
        write_table(ws, log_df)

    # 明細（可篩選）
    ws = wb.create_sheet("本月明細")
    detail = cur.drop(columns="月份")
    last = write_table(ws, detail, formats={"單價": MONEY, "金額": MONEY})
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(detail.columns))}{last}"
    ws.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main(argv=None):
    parser = argparse.ArgumentParser(description="合併多間門市的銷售 Excel 並產出月報")
    parser.add_argument("inputs", nargs="+", type=Path, help="銷售資料檔（.xlsx / .csv），可用萬用字元")
    parser.add_argument("--month", help="報表月份，例如 2026-09；預設為資料中最新的月份")
    parser.add_argument("-o", "--output", type=Path, help="輸出路徑，預設 output/銷售月報_YYYY-MM.xlsx")
    args = parser.parse_args(argv)

    # Windows 的命令列不會展開 *.xlsx，所以自己展開一次
    paths = [Path(p) for pattern in args.inputs for p in (sorted(glob.glob(str(pattern))) or [pattern])]
    raw = pd.concat([load_file(p) for p in paths], ignore_index=True)
    data, log = clean(raw)
    month = args.month or str(data["日期"].max().to_period("M"))
    out = args.output or Path("output") / f"銷售月報_{month}.xlsx"
    build_report(data, log, month, out)

    print(f"讀入 {len(paths)} 個檔案、{len(raw)} 筆原始資料，清理後 {len(data)} 筆")
    for e in log.entries:
        print(f"  [{e['門市']}] {e['問題']}：{e['筆數']} 筆 → {e['處理方式']}")
    print(f"已輸出 {out}")


if __name__ == "__main__":
    main()
