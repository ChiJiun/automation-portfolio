import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import make_sample_data  # noqa: E402
import report  # noqa: E402


def test_normalize_text_merges_fullwidth_and_spaces():
    assert report.normalize_text(" 咖啡豆（半磅） ") == "咖啡豆(半磅)"
    assert report.normalize_text("燕麥  拿鐵") == "燕麥 拿鐵"


def test_standardize_columns_maps_aliases():
    df = pd.DataFrame(columns=["交易日期", "品名", "銷售數量", "售價"])
    assert list(report.standardize_columns(df, "x.xlsx").columns) == ["日期", "商品", "數量", "單價"]


def test_standardize_columns_reports_missing():
    with pytest.raises(ValueError, match="缺少必要欄位：單價"):
        report.standardize_columns(pd.DataFrame(columns=["日期", "商品", "數量"]), "x.xlsx")


def make_raw(rows):
    df = pd.DataFrame(rows, columns=["單號", "日期", "門市", "商品", "類別", "數量", "單價"])
    return df


def test_clean_handles_messy_rows():
    raw = make_raw(
        [
            ("A1", "2026/09/01", "店", "拿鐵", "飲品", "2個", 80),
            ("A1", "2026/09/01", "店", "拿鐵", "飲品", "2個", 80),  # 重複
            ("A2", "2026/13/01", "店", "拿鐵", "飲品", 1, 80),  # 壞日期
            ("A3", "2026-09-02", "店", "拿鐵", "飲品", "N/A", 80),  # 壞數量
            ("A4", "2026-09-02", "店", " 拿鐵 ", "飲品", 3, None),  # 缺單價
        ]
    )
    data, log = report.clean(raw)
    assert list(data["單號"]) == ["A1", "A4"]
    assert data.loc[data["單號"] == "A4", "單價"].item() == 80  # 用同商品中位數補
    assert data["金額"].tolist() == [160, 240]
    problems = {e["問題"]: e["筆數"] for e in log.entries}
    assert problems == {"日期無法辨識": 1, "數量無法辨識或為 0": 1, "單價空白": 1, "完全重複的列": 1}


def test_end_to_end_report(tmp_path):
    make_sample_data.main(tmp_path / "data")
    out = tmp_path / "report.xlsx"
    report.main([str(tmp_path / "data" / "*.xlsx"), "--month", "2026-08", "-o", str(out)])

    wb = load_workbook(out)
    assert wb.sheetnames == ["摘要", "門市比較", "商品排行", "每日趨勢", "清理紀錄", "本月明細"]
    assert wb["摘要"]["A1"].value == "2026 年 8 月 銷售月報"
    stores = {row[0] for row in wb["門市比較"].iter_rows(min_row=2, values_only=True)}
    assert stores == {"台北店", "台中店", "高雄店"}
    # 全形括號版本要和半形版本合併成同一個商品
    items = [row[1] for row in wb["商品排行"].iter_rows(min_row=2, values_only=True)]
    assert len(items) == len(set(items)) == len(make_sample_data.PRODUCTS)
    assert wb["每日趨勢"].max_row == 32  # 標題 + 8 月 31 天


def test_missing_month_raises(tmp_path):
    make_sample_data.main(tmp_path / "data")
    with pytest.raises(ValueError, match="沒有 2025-01"):
        report.main([str(tmp_path / "data" / "*.xlsx"), "--month", "2025-01", "-o", str(tmp_path / "x.xlsx")])
